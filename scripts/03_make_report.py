#!/usr/bin/env python3
"""
03_make_report.py  v2
BCR Repertoire — 單樣本 Excel 報告
用法：python3 03_make_report.py --airr <tsv> --chain <IGH|IGK|IGL>
      --sample <name> --stats <json> --output <xlsx>
"""

import argparse, json
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Style ─────────────────────────────────────────────────────────────────────
HDR_FILL   = PatternFill("solid", fgColor="1F4E79")
HDR_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=13, color="1F4E79")
BODY_FONT  = Font(name="Arial", size=10)
ALT_FILL   = PatternFill("solid", fgColor="DCE6F1")
THIN       = Side(style="thin", color="AAAAAA")
BORDER     = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def hdr(cell, fill=HDR_FILL):
    cell.fill = fill; cell.font = HDR_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

def body(cell, alt=False):
    cell.font = BODY_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    cell.border = BORDER
    if alt: cell.fill = ALT_FILL

def write_table(ws, headers, rows, start_row=1, start_col=1, pct_cols=None):
    pct_cols = pct_cols or []
    for ci, h in enumerate(headers, start_col):
        hdr(ws.cell(row=start_row, column=ci, value=h))
    for ri, row in enumerate(rows, start_row + 1):
        alt = (ri - start_row) % 2 == 0
        for ci, val in enumerate(row, start_col):
            c = ws.cell(row=ri, column=ci, value=val)
            body(c, alt)
            if ci in pct_cols and isinstance(val, (int, float)):
                c.number_format = "0.00%"
            elif isinstance(val, float):
                c.number_format = "0.00"
    return ri if rows else start_row

def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def title_cell(ws, text, span, row=1):
    ws.merge_cells(f"A{row}:{span}{row}")
    c = ws[f"A{row}"]
    c.value = text; c.font = TITLE_FONT
    c.alignment = Alignment(horizontal="left")
    ws.row_dimensions[row].height = 22

# ── Helpers ───────────────────────────────────────────────────────────────────
def safe_float(val, default=None):
    try:
        v = float(val)
        return v if not np.isnan(v) else default
    except (TypeError, ValueError):
        return default

def parse_first_call(val):
    if pd.isna(val): return ""
    return str(val).split(",")[0].strip()

def shannon(counts):
    counts = counts[counts > 0]
    if len(counts) == 0: return 0.0, 0.0
    total = counts.sum()
    p = counts / total
    H = float(-np.sum(p * np.log2(p + 1e-12)))
    N = len(counts)
    return round(H, 4), round(float(1 - H / np.log2(N)) if N > 1 else 0.0, 4)

def get_productive(df):
    if "productive" in df.columns and df["productive"].notna().sum() > 0:
        prod = df[df["productive"].str.upper() == "T"].copy()
        if len(prod) > 0: return prod
    if "stop_codon" in df.columns and "vj_in_frame" in df.columns:
        prod = df[(df["stop_codon"] == "F") & (df["vj_in_frame"] == "T")].copy()
        if len(prod) > 0: return prod
    if "stop_codon" in df.columns:
        prod = df[df["stop_codon"] == "F"].copy()
        if len(prod) > 0: return prod
    return df.copy()

def parse_gene_calls(series, top_n=50):
    first = series.dropna().apply(lambda x: str(x).split(",")[0].strip())
    counts = first.value_counts()
    return counts.head(top_n), counts.sum()

# ── Build clonotype table ─────────────────────────────────────────────────────
def build_clones(df_prod, chain):
    """
    Group by junction_aa (CDR3 AA). For each clone, aggregate:
    read count, CDR3_NT, V/D/J genes, SHM%, CDR1/2, np1/2, complete_vdj.
    Returns sorted DataFrame.
    """
    has_cdr3 = "junction_aa" in df_prod.columns and df_prod["junction_aa"].notna().sum() > 0

    if not has_cdr3:
        # No CDR3: treat each unique sequence_id as a clone
        grp = df_prod.groupby("sequence_id")
    else:
        grp = df_prod.groupby("junction_aa")

    rows = []
    total_reads = len(df_prod)

    for clone_id, members in grp:
        read_count = len(members)
        freq = read_count / total_reads if total_reads > 0 else 0

        # Representative = row with highest v_identity (or first)
        if "v_identity" in members.columns and members["v_identity"].notna().sum() > 0:
            rep = members.loc[members["v_identity"].idxmax()]
        else:
            rep = members.iloc[0]

        cdr3_aa  = clone_id if has_cdr3 else ""
        cdr3_len = len(cdr3_aa) - 2 if has_cdr3 and len(cdr3_aa) > 2 else ""

        # CDR3 nucleotide — use cdr3 column if available, else junction minus anchors
        cdr3_nt = ""
        if "cdr3" in members.columns and pd.notna(rep.get("cdr3")):
            cdr3_nt = str(rep["cdr3"])
        elif "junction" in members.columns and pd.notna(rep.get("junction")):
            j = str(rep["junction"])
            cdr3_nt = j[3:-3] if len(j) > 6 else j

        v_gene = parse_first_call(rep.get("v_call"))
        d_gene = parse_first_call(rep.get("d_call")) if chain == "IGH" else ""
        j_gene = parse_first_call(rep.get("j_call"))

        v_id = safe_float(rep.get("v_identity"))
        d_id = safe_float(rep.get("d_identity"))
        j_id = safe_float(rep.get("j_identity"))
        v_shm = round(100 - v_id, 3) if v_id is not None else ""
        d_shm = round(100 - d_id, 3) if d_id is not None else ""
        j_shm = round(100 - j_id, 3) if j_id is not None else ""

        cdr1 = rep.get("cdr1_aa", "") if "cdr1_aa" in members.columns else ""
        cdr2 = rep.get("cdr2_aa", "") if "cdr2_aa" in members.columns else ""
        cdr1 = "" if pd.isna(cdr1) else str(cdr1)
        cdr2 = "" if pd.isna(cdr2) else str(cdr2)

        np1_mean = ""
        np2_mean = ""
        if "np1_length" in members.columns:
            np1_vals = members["np1_length"].dropna()
            np1_mean = round(float(np1_vals.mean()), 1) if len(np1_vals) > 0 else ""
        if "np2_length" in members.columns:
            np2_vals = members["np2_length"].dropna()
            np2_mean = round(float(np2_vals.mean()), 1) if len(np2_vals) > 0 else ""

        complete_vdj = ""
        if "complete_vdj" in members.columns:
            cvdj = members["complete_vdj"].dropna()
            if len(cvdj) > 0:
                pct = cvdj.apply(lambda x: str(x).upper() == "T").mean() * 100
                complete_vdj = round(float(pct), 1)

        rows.append({
            "Read_Count":       read_count,
            "Frequency_%":      round(freq, 6),
            "CDR3_AA":          cdr3_aa,
            "CDR3_AA_length":   cdr3_len,
            "CDR3_NT":          cdr3_nt,
            "V_gene":           v_gene,
            "D_gene":           d_gene,
            "J_gene":           j_gene,
            "V_SHM_%":          v_shm,
            "D_SHM_%":          d_shm,
            "J_SHM_%":          j_shm,
            "CDR1_AA":          cdr1,
            "CDR2_AA":          cdr2,
            "np1_length_mean":  np1_mean,
            "np2_length_mean":  np2_mean,
            "complete_vdj":     complete_vdj,
        })

    clone_df = pd.DataFrame(rows).sort_values("Read_Count", ascending=False).reset_index(drop=True)
    clone_df.insert(0, "Rank", range(1, len(clone_df) + 1))
    return clone_df

# ── Main ──────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--airr",   required=True)
    parser.add_argument("--chain",  required=True, choices=["IGH","IGK","IGL"])
    parser.add_argument("--sample", required=True)
    parser.add_argument("--stats",  required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    df = pd.read_csv(args.airr, sep="\t", low_memory=False)
    with open(args.stats) as f:
        stats = json.load(f)

    df_prod = get_productive(df)
    total   = len(df)
    n_prod  = len(df_prod)

    has_cdr3 = "junction_aa" in df.columns and df["junction_aa"].notna().sum() > 0

    clone_df  = build_clones(df_prod, args.chain)
    n_clones  = len(clone_df)
    clone_counts = clone_df["Read_Count"]
    H, clonality = shannon(clone_counts)

    wb = Workbook()
    wb.remove(wb.active)

    # ══ Sheet 1: Summary ══════════════════════════════════════════════════════
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    title_cell(ws, f"BCR Repertoire Report — {args.sample}", "C")

    platform     = stats.get("platform", "")
    reads_raw    = stats.get("reads_raw", "N/A")
    reads_merged = stats.get("reads_merged", stats.get("reads_qc", "N/A"))
    reads_trim   = stats.get("reads_trimmed", "N/A")
    reads_ib     = stats.get("reads_igblast", total)

    v_id_all  = safe_float(df["v_identity"].mean())   if "v_identity" in df.columns else None
    v_id_prod = safe_float(df_prod["v_identity"].mean()) if "v_identity" in df_prod.columns and n_prod > 0 else None

    rows_sum = [
        ("Platform",                   platform,  ""),
        ("Chain",                      args.chain,""),
        ("── Read Processing ──",      "",        ""),
        ("Raw reads",                  reads_raw, ""),
        ("After QC / merge",           reads_merged,""),
        ("After adapter trim",         reads_trim,""),
        ("IgBLAST input",              reads_ib,  ""),
        ("── IgBLAST Results ──",      "",        ""),
        ("Total sequences",            total,     ""),
        ("Productive sequences",       n_prod,    f"{n_prod/total*100:.1f}%" if total>0 else ""),
        ("Non-productive",             total-n_prod, f"{(total-n_prod)/total*100:.1f}%" if total>0 else ""),
        ("Has V call",                 int(df["v_call"].notna().sum()),""),
        ("Has J call",                 int(df["j_call"].notna().sum()),""),
        ("Has CDR3 (junction_aa)",     int(df["junction_aa"].notna().sum()) if "junction_aa" in df.columns else 0,""),
        ("── Repertoire (Productive) ──","",""),
        ("Unique clones",              n_clones,  ""),
        ("Singleton clones",           int((clone_counts==1).sum()), f"{(clone_counts==1).sum()/n_clones*100:.1f}%" if n_clones>0 else ""),
        ("Max clone size",             int(clone_counts.max()) if n_clones>0 else 0,""),
        ("Shannon entropy H",          H,         ""),
        ("Clonality (1-H/log2N)",      clonality, ""),
        ("Mean V identity % (all)",    round(v_id_all,2)  if v_id_all  is not None else "N/A",""),
        ("Mean V identity % (prod)",   round(v_id_prod,2) if v_id_prod is not None else "N/A",""),
        ("Mean V SHM % (all)",         round(100-v_id_all,2)  if v_id_all  is not None else "N/A",""),
        ("Mean V SHM % (prod)",        round(100-v_id_prod,2) if v_id_prod is not None else "N/A",""),
    ]
    write_table(ws, ["Metric","Value","Percent"], rows_sum, start_row=3)
    set_widths(ws, {"A":38,"B":16,"C":12})

    # ══ Sheet 2: Clones (ALL) ═════════════════════════════════════════════════
    ws2 = wb.create_sheet("Clones")
    ws2.sheet_view.showGridLines = False

    chain_label = {"IGH":"IgH","IGK":"Igκ","IGL":"Igλ"}.get(args.chain, args.chain)
    cdr3_def    = "相同 CDR3 AA 序列" if has_cdr3 else "sequence_id"
    ws2.merge_cells("A1:Q1")
    t = ws2["A1"]
    t.value = f"{chain_label} Clonotypes — 定義: {cdr3_def}  |  共 {n_clones} clonotypes"
    t.font = TITLE_FONT
    t.alignment = Alignment(horizontal="left")
    ws2.row_dimensions[1].height = 20

    clone_headers = ["Rank","Read_Count","Frequency_%","CDR3_AA","CDR3_AA_length",
                     "CDR3_NT","V_gene","D_gene","J_gene","V_SHM_%","D_SHM_%","J_SHM_%",
                     "CDR1_AA","CDR2_AA","np1_length_mean","np2_length_mean","complete_vdj"]

    clone_rows = [tuple(row[h] for h in clone_headers) for _, row in clone_df.iterrows()]
    write_table(ws2, clone_headers, clone_rows, start_row=2, pct_cols=[3])

    clone_widths = {"A":6,"B":12,"C":12,"D":22,"E":12,
                    "F":36,"G":18,"H":22,"I":14,"J":10,
                    "K":10,"L":10,"M":14,"N":14,"O":14,"P":14,"Q":14}
    set_widths(ws2, clone_widths)

    # ══ Sheet 3: CDR3 Length ══════════════════════════════════════════════════
    ws3 = wb.create_sheet("CDR3_Length")
    ws3.sheet_view.showGridLines = False
    title_cell(ws3, f"CDR3 Length Distribution (Productive) — {args.sample}", "D")

    if has_cdr3 and n_clones > 0:
        cdr3_lens = clone_df["CDR3_AA_length"].dropna().astype(int)
        len_range = range(int(cdr3_lens.min()), int(cdr3_lens.max())+1)
        total_reads_prod = clone_df["Read_Count"].sum()
        rows_cdr3 = []
        for l in len_range:
            mask = cdr3_lens == l
            cl_n = int(mask.sum())
            rd_n = int(clone_df.loc[mask.values, "Read_Count"].sum())
            rows_cdr3.append((l, cl_n, rd_n, round(rd_n/total_reads_prod,6) if total_reads_prod>0 else 0))
        write_table(ws3, ["CDR3 AA Length","Clone Count","Read Count","Frequency (%)"],
                    rows_cdr3, start_row=3, pct_cols=[4])
    else:
        ws3.cell(row=3,column=1,value="N/A — no CDR3 data")
    set_widths(ws3, {"A":16,"B":14,"C":12,"D":14})

    # ══ Sheet 4: V Usage ══════════════════════════════════════════════════════
    ws4 = wb.create_sheet("V_Usage")
    ws4.sheet_view.showGridLines = False
    title_cell(ws4, f"V Gene Usage (Productive) — {args.sample}", "E")

    v_clone = clone_df.groupby("V_gene")["Read_Count"].agg(["count","sum"]).reset_index()
    v_clone.columns = ["V_gene","Clone_Count","Read_Count"]
    v_clone = v_clone.sort_values("Clone_Count", ascending=False).reset_index(drop=True)
    total_clones_v = v_clone["Clone_Count"].sum()
    rows_v = [(i+1, row["V_gene"], int(row["Clone_Count"]), int(row["Read_Count"]),
               round(row["Clone_Count"]/total_clones_v,6) if total_clones_v>0 else 0)
              for i, row in v_clone.iterrows()]
    write_table(ws4, ["Rank","V_gene","Clone Count","Read Count","Frequency (%)"],
                rows_v, start_row=3, pct_cols=[5])
    set_widths(ws4, {"A":6,"B":22,"C":14,"D":12,"E":14})

    # ══ Sheet 5: D Usage (IGH only) ═══════════════════════════════════════════
    if args.chain == "IGH":
        ws5 = wb.create_sheet("D_Usage")
        ws5.sheet_view.showGridLines = False
        title_cell(ws5, f"D Gene Usage (Productive) — {args.sample}", "E")
        d_clone = clone_df[clone_df["D_gene"] != ""].groupby("D_gene")["Read_Count"].agg(["count","sum"]).reset_index()
        d_clone.columns = ["D_gene","Clone_Count","Read_Count"]
        d_clone = d_clone.sort_values("Clone_Count", ascending=False).reset_index(drop=True)
        total_clones_d = d_clone["Clone_Count"].sum()
        rows_d = [(i+1, row["D_gene"], int(row["Clone_Count"]), int(row["Read_Count"]),
                   round(row["Clone_Count"]/total_clones_d,6) if total_clones_d>0 else 0)
                  for i, row in d_clone.iterrows()]
        write_table(ws5, ["Rank","D_gene","Clone Count","Read Count","Frequency (%)"],
                    rows_d, start_row=3, pct_cols=[5])
        set_widths(ws5, {"A":6,"B":22,"C":14,"D":12,"E":14})

    # ══ Sheet 6: J Usage ══════════════════════════════════════════════════════
    ws6 = wb.create_sheet("J_Usage")
    ws6.sheet_view.showGridLines = False
    title_cell(ws6, f"J Gene Usage (Productive) — {args.sample}", "E")
    j_clone = clone_df[clone_df["J_gene"] != ""].groupby("J_gene")["Read_Count"].agg(["count","sum"]).reset_index()
    j_clone.columns = ["J_gene","Clone_Count","Read_Count"]
    j_clone = j_clone.sort_values("Clone_Count", ascending=False).reset_index(drop=True)
    total_clones_j = j_clone["Clone_Count"].sum()
    rows_j = [(i+1, row["J_gene"], int(row["Clone_Count"]), int(row["Read_Count"]),
               round(row["Clone_Count"]/total_clones_j,6) if total_clones_j>0 else 0)
              for i, row in j_clone.iterrows()]
    write_table(ws6, ["Rank","J_gene","Clone Count","Read Count","Frequency (%)"],
                rows_j, start_row=3, pct_cols=[5])
    set_widths(ws6, {"A":6,"B":22,"C":14,"D":12,"E":14})

    # ══ Sheet 7: SHM Distribution ═════════════════════════════════════════════
    ws7 = wb.create_sheet("SHM_Distribution")
    ws7.sheet_view.showGridLines = False
    title_cell(ws7, f"V Gene SHM Distribution (Productive) — {args.sample}", "D")

    v_shm_vals = clone_df["V_SHM_%"].apply(lambda x: safe_float(x)).dropna()
    bins   = list(range(0, 21)) + [float("inf")]
    labels = [f"{i}–{i+1}" for i in range(20)] + ["≥20"]
    total_c = len(v_shm_vals)
    rows_shm = []
    for lo, hi, lab in zip(bins, bins[1:], labels):
        if hi == float("inf"):
            mask = v_shm_vals >= lo
        else:
            mask = (v_shm_vals >= lo) & (v_shm_vals < hi)
        n = int(mask.sum())
        # read count for these clones
        rd = int(clone_df.loc[v_shm_vals[mask].index, "Read_Count"].sum()) if n > 0 else 0
        rows_shm.append((lab, n, rd, round(n/total_c,6) if total_c>0 else 0))
    write_table(ws7, ["V SHM (%) Bin","Clone Count","Read Count","Frequency (%)"],
                rows_shm, start_row=3, pct_cols=[4])
    set_widths(ws7, {"A":14,"B":14,"C":12,"D":14})

    wb.save(args.output)
    print(f"✓ Report saved: {args.output}  ({n_clones} clones)")

if __name__ == "__main__":
    main()
